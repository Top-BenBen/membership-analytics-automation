"""Reusable membership analytics and Excel reporting automation.

This public portfolio version is industry-neutral and uses generic plan labels.
It expects an active-membership snapshot, not a full historical member register.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from datetime import date
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = [
    "MEMBERSHIP NO", "Date of Birth", "GENDER", "COMPANY", "INDUSTRY",
    "Registration Date", "Join Date", "BENEFIT OPTION", "MEMBER TYPE",
    "RELATIONSHIP TYPE",
]
NAVY, BLUE, LIGHT_BLUE, WHITE, GREY = "17365D", "2F75B5", "D9EAF7", "FFFFFF", "E7E6E6"


def normalise_text(series: pd.Series) -> pd.Series:
    return (series.astype("string").str.strip().str.replace(r"\s+", " ", regex=True)
            .replace({"": pd.NA, "nan": pd.NA, "None": pd.NA}))


def age_years(dob: pd.Series, as_of: pd.Timestamp) -> pd.Series:
    dob = pd.to_datetime(dob, errors="coerce")
    ages = as_of.year - dob.dt.year - ((as_of.month < dob.dt.month) | ((as_of.month == dob.dt.month) & (as_of.day < dob.dt.day)))
    return ages.where(dob.notna() & (dob <= as_of) & (ages >= 0))


def age_band(age):
    if pd.isna(age):
        return "UNKNOWN"
    low = (int(age) // 5) * 5
    return "65+" if low >= 65 else f"{low}–{low + 4}"


def tenure_band(months):
    if pd.isna(months): return "UNKNOWN"
    if months < 3: return "<3 months"
    if months < 6: return "3–5 months"
    if months < 9: return "6–8 months"
    if months < 12: return "9–11 months"
    return "12+ months"


def load_and_clean(path: Path, as_of: pd.Timestamp) -> tuple[pd.DataFrame, pd.DataFrame]:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        raw = pd.read_excel(path)
    elif path.suffix.lower() == ".csv":
        raw = pd.read_csv(path)
    else:
        raise ValueError("Input must be a CSV or Excel file.")

    raw.columns = [str(c).strip() for c in raw.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in raw.columns]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))

    df = raw[REQUIRED_COLUMNS].copy()
    for col in ["MEMBERSHIP NO", "GENDER", "COMPANY", "INDUSTRY", "BENEFIT OPTION", "MEMBER TYPE", "RELATIONSHIP TYPE"]:
        df[col] = normalise_text(df[col])
    for col in ["GENDER", "COMPANY", "INDUSTRY", "BENEFIT OPTION", "MEMBER TYPE", "RELATIONSHIP TYPE"]:
        df[col] = df[col].str.upper()
    for col in ["Date of Birth", "Registration Date", "Join Date"]:
        df[col] = pd.to_datetime(df[col], errors="coerce")

    checks = []
    checks.append(("Input rows", len(raw), "Rows read from source file"))
    checks.append(("Rows after exact duplicate removal", len(df.drop_duplicates()), "Exact duplicate records removed"))
    before = len(df)
    df = df.drop_duplicates()
    checks.append(("Rows with missing membership number", int(df["MEMBERSHIP NO"].isna().sum()), "Review before publishing results"))
    checks.append(("Duplicate membership numbers", int(df["MEMBERSHIP NO"].duplicated(keep=False).sum()), "May be legitimate only if source structure allows it; investigate"))
    for col in REQUIRED_COLUMNS:
        checks.append((f"Missing: {col}", int(df[col].isna().sum()), "Missing-value count after standardisation"))

    df["AGE"] = age_years(df["Date of Birth"], as_of)
    df["AGE COHORT"] = df["AGE"].apply(age_band)
    df["TENURE MONTHS"] = ((as_of.year - df["Join Date"].dt.year) * 12 + as_of.month - df["Join Date"].dt.month).where(df["Join Date"].notna())
    df.loc[df["TENURE MONTHS"] < 0, "TENURE MONTHS"] = np.nan
    df["TENURE BAND"] = df["TENURE MONTHS"].apply(tenure_band)
    quality = pd.DataFrame(checks, columns=["Check", "Count", "Interpretation"])
    return df, quality


def count_table(df: pd.DataFrame, col: str, label: str = "Category") -> pd.DataFrame:
    out = df[col].fillna("UNKNOWN").value_counts(dropna=False).rename_axis(label).reset_index(name="Active Members")
    out["Share of Population"] = out["Active Members"] / max(len(df), 1)
    return out


def build_report(df: pd.DataFrame, quality: pd.DataFrame, output: Path, as_of: pd.Timestamp) -> None:
    total = len(df)
    unique_members = df["MEMBERSHIP NO"].nunique(dropna=True)
    valid_ages = df["AGE"].dropna()
    summary = pd.DataFrame([
        ("Reporting date", as_of.date().isoformat()),
        ("Rows in cleaned snapshot", total),
        ("Unique membership numbers", unique_members),
        ("Companies represented", df["COMPANY"].nunique(dropna=True)),
        ("Industries represented", df["INDUSTRY"].nunique(dropna=True)),
        ("Average age", round(valid_ages.mean(), 1) if not valid_ages.empty else None),
        ("Median age", round(valid_ages.median(), 1) if not valid_ages.empty else None),
        ("Female members (%)", round((df["GENDER"].eq("FEMALE").sum() / total) * 100, 1) if total else 0),
        ("Principal members", int(df["MEMBER TYPE"].eq("PRINCIPAL").sum())),
        ("Dependent members", int(df["MEMBER TYPE"].eq("DEPENDANT").sum())),
    ], columns=["Metric", "Value"])

    sheets = {
        "Executive Summary": summary,
        "Demographics": pd.concat([count_table(df, "GENDER", "Gender"), count_table(df, "RELATIONSHIP TYPE", "Relationship Type")], ignore_index=True),
        "By Company": count_table(df, "COMPANY", "Company"),
        "By Industry": count_table(df, "INDUSTRY", "Industry"),
        "Plan Analysis": count_table(df, "BENEFIT OPTION", "Benefit Option"),
        "Age Analysis": count_table(df, "AGE COHORT", "Age Cohort"),
        "Tenure": count_table(df, "TENURE BAND", "Tenure Band"),
        "Data Quality": quality,
        "Cleaned Data": df,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, table in sheets.items():
            table.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(output)
    thin = Side(style="thin", color="D9E1F2")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            lengths = [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, min(ws.max_row, 300) + 1)]
            ws.column_dimensions[letter].width = min(max(max(lengths, default=10) + 2, 12), 38)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                if cell.data_type == "n": cell.number_format = "#,##0.0" if isinstance(cell.value, float) else "#,##0"
        if ws.title in {"By Company", "By Industry", "Plan Analysis", "Age Analysis", "Tenure"} and ws.max_row > 2:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = ws.title
            chart.y_axis.title = ws.cell(1, 1).value
            chart.x_axis.title = "Active Members"
            data = Reference(ws, min_col=2, min_row=1, max_row=min(ws.max_row, 16))
            cats = Reference(ws, min_col=1, min_row=2, max_row=min(ws.max_row, 16))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height, chart.width = 7, 13
            ws.add_chart(chart, "E2")
    wb.save(output)


def main():
    parser = argparse.ArgumentParser(description="Clean an active-membership snapshot and create an Excel analytics report.")
    parser.add_argument("input", nargs="?", default="data/sample/sample_membership.csv", help="Input CSV or Excel file")
    parser.add_argument("output", nargs="?", default="outputs/membership_analytics_report.xlsx", help="Output Excel workbook")
    parser.add_argument("--as-of", default=date.today().isoformat(), help="Reporting date YYYY-MM-DD")
    args = parser.parse_args()
    as_of = pd.Timestamp(args.as_of)
    cleaned, quality = load_and_clean(Path(args.input), as_of)
    build_report(cleaned, quality, Path(args.output), as_of)
    print(f"Report created: {Path(args.output).resolve()}")
    print(f"Rows analysed: {len(cleaned):,}; unique membership numbers: {cleaned['MEMBERSHIP NO'].nunique(dropna=True):,}")

if __name__ == "__main__":
    main()
